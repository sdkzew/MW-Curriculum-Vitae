from pathlib import Path
from zipfile import BadZipFile, ZipFile

import docx
import pdfplumber


MAX_PDF_PAGES = 50
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 500_000


class DocumentExtractionError(ValueError):
    """Raised when an uploaded document is invalid or unsafe to process."""


def extract_text(filepath: str) -> str:
    extension = Path(filepath).suffix.lower()
    try:
        if extension == ".pdf":
            text = _extract_pdf(filepath)
        elif extension == ".docx":
            text = _extract_docx(filepath)
        else:
            raise DocumentExtractionError("Unsupported document type")
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError("Document could not be parsed") from error

    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentExtractionError("Document contains too much text")
    return text


def _extract_pdf(filepath: str) -> str:
    text = []
    with pdfplumber.open(filepath) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise DocumentExtractionError(
                f"PDF has more than the supported {MAX_PDF_PAGES} pages"
            )
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
    return "\n".join(text)


def _extract_docx(filepath: str) -> str:
    try:
        with ZipFile(filepath) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentExtractionError("Invalid DOCX structure")
            uncompressed_size = sum(info.file_size for info in archive.infolist())
            if uncompressed_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentExtractionError("DOCX expands beyond the supported size")
    except BadZipFile as error:
        raise DocumentExtractionError("Invalid DOCX archive") from error

    document = docx.Document(filepath)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    parts.append(value)
    return "\n".join(parts)
