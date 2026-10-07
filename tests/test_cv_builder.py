import unittest
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from PIL import Image
from werkzeug.datastructures import MultiDict
from werkzeug.datastructures import FileStorage

from analyzer.cv_builder import (
    CVBuildError,
    build_cv_document,
    prepare_profile_photo,
    profile_from_form,
)


class CVBuilderTests(unittest.TestCase):
    def sample_form(self, template="europass"):
        return MultiDict(
            [
                ("template", template),
                ("full_name", "Alex Morgan"),
                ("headline", "Product Designer"),
                ("email", "alex@example.com"),
                ("location", "Bucharest, Romania"),
                ("summary", "Designer focused on useful digital products."),
                ("experience_role", "Senior Product Designer"),
                ("experience_organization", "MORPH WRLD"),
                ("experience_location", "Remote"),
                ("experience_dates", "2023 — Present"),
                ("experience_details", "Improved conversion by 18%.\nBuilt a design system."),
                ("education_qualification", "BA Design"),
                ("education_institution", "Design University"),
                ("education_location", "Bucharest"),
                ("education_dates", "2019 — 2022"),
                ("education_details", "Interaction design."),
                ("skills", "Research · Prototyping"),
                ("digital_skills", "Figma · Jira"),
                ("languages", "Romanian — native\nEnglish — C1"),
            ]
        )

    def document_text(self, document):
        paragraph_text = [paragraph.text for paragraph in document.paragraphs]
        table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(paragraph_text + table_text)

    def test_profile_parses_repeatable_sections(self):
        profile, template = profile_from_form(self.sample_form())

        self.assertEqual(template, "europass")
        self.assertEqual(profile.full_name, "Alex Morgan")
        self.assertEqual(profile.experiences[0].organization, "MORPH WRLD")
        self.assertEqual(profile.education[0].qualification, "BA Design")

    def test_all_templates_create_valid_docx_documents(self):
        for template in ("europass", "ats", "modern"):
            with self.subTest(template=template):
                profile, _ = profile_from_form(self.sample_form(template))
                output = build_cv_document(profile, template)
                document = Document(output)
                text = self.document_text(document)

                self.assertIn("Alex Morgan", text)
                self.assertIn("Senior Product Designer", text)
                self.assertIn("Research", text)

    def test_missing_name_is_rejected(self):
        with self.assertRaises(CVBuildError):
            profile_from_form({"template": "ats", "full_name": ""})

    def test_unknown_template_is_rejected(self):
        with self.assertRaises(CVBuildError):
            profile_from_form({"template": "official", "full_name": "Alex"})

    def test_photo_is_safely_processed_and_embedded(self):
        source = BytesIO()
        Image.new("RGB", (900, 600), (92, 64, 170)).save(source, format="JPEG")
        source.seek(0)
        upload = FileStorage(stream=source, filename="portrait.jpg", content_type="image/jpeg")
        photo = prepare_profile_photo(upload, "en")
        profile, template = profile_from_form(self.sample_form("modern"), "en")

        output = build_cv_document(profile, template, photo, "en")
        payload = output.getvalue()
        with ZipFile(BytesIO(payload)) as archive:
            self.assertTrue(any(name.startswith("word/media/") for name in archive.namelist()))
        document = Document(BytesIO(payload))
        text = self.document_text(document)
        self.assertIn("WORK EXPERIENCE", text)

    def test_invalid_photo_is_rejected(self):
        upload = FileStorage(stream=BytesIO(b"not an image"), filename="portrait.png")
        with self.assertRaisesRegex(CVBuildError, "valid JPG or PNG"):
            prepare_profile_photo(upload, "en")


if __name__ == "__main__":
    unittest.main()
