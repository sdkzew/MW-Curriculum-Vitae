from dataclasses import dataclass, field
from io import BytesIO
from itertools import zip_longest
from typing import Mapping

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TEMPLATES = {"europass", "ats", "modern"}


class CVBuildError(ValueError):
    """Raised when CV builder input is missing or exceeds safe limits."""


@dataclass
class Experience:
    role: str
    organization: str = ""
    location: str = ""
    dates: str = ""
    details: str = ""


@dataclass
class Education:
    qualification: str
    institution: str = ""
    location: str = ""
    dates: str = ""
    details: str = ""


@dataclass
class CVProfile:
    full_name: str
    headline: str = ""
    email: str = ""
    phone: str = ""
    location: str = ""
    linkedin: str = ""
    website: str = ""
    summary: str = ""
    skills: str = ""
    digital_skills: str = ""
    languages: str = ""
    certifications: str = ""
    projects: str = ""
    additional: str = ""
    experiences: list[Experience] = field(default_factory=list)
    education: list[Education] = field(default_factory=list)


def _clean(value: object, limit: int, label: str) -> str:
    text = str(value or "").replace("\x00", "").strip()
    if len(text) > limit:
        raise CVBuildError(f"{label} depășește limita de {limit} caractere.")
    return text


def _values(form: Mapping, key: str) -> list[str]:
    if hasattr(form, "getlist"):
        return list(form.getlist(key))
    value = form.get(key, [])
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]


def profile_from_form(form: Mapping) -> tuple[CVProfile, str]:
    template = _clean(form.get("template", "ats"), 20, "Template")
    if template not in TEMPLATES:
        raise CVBuildError("Template necunoscut.")

    full_name = _clean(form.get("full_name"), 120, "Numele")
    if not full_name:
        raise CVBuildError("Numele complet este obligatoriu.")

    experience_rows = zip_longest(
        _values(form, "experience_role"),
        _values(form, "experience_organization"),
        _values(form, "experience_location"),
        _values(form, "experience_dates"),
        _values(form, "experience_details"),
        fillvalue="",
    )
    experiences = []
    for role, organization, location, dates, details in list(experience_rows)[:8]:
        role = _clean(role, 180, "Rolul")
        if not role:
            continue
        experiences.append(
            Experience(
                role=role,
                organization=_clean(organization, 180, "Organizația"),
                location=_clean(location, 120, "Locația"),
                dates=_clean(dates, 100, "Perioada"),
                details=_clean(details, 4_000, "Descrierea experienței"),
            )
        )

    education_rows = zip_longest(
        _values(form, "education_qualification"),
        _values(form, "education_institution"),
        _values(form, "education_location"),
        _values(form, "education_dates"),
        _values(form, "education_details"),
        fillvalue="",
    )
    education = []
    for qualification, institution, location, dates, details in list(education_rows)[:8]:
        qualification = _clean(qualification, 180, "Calificarea")
        if not qualification:
            continue
        education.append(
            Education(
                qualification=qualification,
                institution=_clean(institution, 180, "Instituția"),
                location=_clean(location, 120, "Locația studiilor"),
                dates=_clean(dates, 100, "Perioada studiilor"),
                details=_clean(details, 3_000, "Descrierea studiilor"),
            )
        )

    profile = CVProfile(
        full_name=full_name,
        headline=_clean(form.get("headline"), 180, "Titlul profesional"),
        email=_clean(form.get("email"), 254, "Email"),
        phone=_clean(form.get("phone"), 80, "Telefon"),
        location=_clean(form.get("location"), 160, "Locație"),
        linkedin=_clean(form.get("linkedin"), 300, "LinkedIn"),
        website=_clean(form.get("website"), 300, "Website"),
        summary=_clean(form.get("summary"), 4_000, "Profil profesional"),
        skills=_clean(form.get("skills"), 2_000, "Competențe"),
        digital_skills=_clean(form.get("digital_skills"), 2_000, "Competențe digitale"),
        languages=_clean(form.get("languages"), 2_000, "Limbi"),
        certifications=_clean(form.get("certifications"), 2_000, "Certificări"),
        projects=_clean(form.get("projects"), 3_000, "Proiecte"),
        additional=_clean(form.get("additional"), 3_000, "Informații suplimentare"),
        experiences=experiences,
        education=education,
    )
    return profile, template


def _set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _add_rule(paragraph, color: str, size: int = 8) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    properties.append(borders)


def _add_section_heading(document: Document, title: str, color: RGBColor, template: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(11)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(10 if template == "europass" else 9.5)
    run.font.color.rgb = color
    run.font.name = "Aptos Display"
    if template != "ats":
        _add_rule(paragraph, str(color), 7)


def _add_contact_line(document: Document, profile: CVProfile, color: RGBColor) -> None:
    contact = [profile.email, profile.phone, profile.location, profile.linkedin, profile.website]
    contact = [item for item in contact if item]
    if not contact:
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run("  •  ".join(contact))
    run.font.size = Pt(8.5)
    run.font.color.rgb = color


def _add_bullets(document: Document, text: str) -> None:
    for line in (line.strip(" •\t-") for line in text.splitlines()):
        if not line:
            continue
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.left_indent = Cm(.45)
        paragraph.add_run(line)


def _add_experience(document: Document, item: Experience, accent: RGBColor) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    role = paragraph.add_run(item.role)
    role.bold = True
    role.font.size = Pt(10.5)
    role.font.color.rgb = accent
    if item.organization:
        paragraph.add_run(f"  |  {item.organization}").bold = True
    meta = " · ".join(value for value in (item.dates, item.location) if value)
    if meta:
        meta_paragraph = document.add_paragraph(meta)
        meta_paragraph.paragraph_format.space_after = Pt(3)
        meta_paragraph.runs[0].italic = True
        meta_paragraph.runs[0].font.size = Pt(8.5)
        meta_paragraph.runs[0].font.color.rgb = RGBColor(95, 99, 110)
    _add_bullets(document, item.details)


def _add_education(document: Document, item: Education, accent: RGBColor) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    qualification = paragraph.add_run(item.qualification)
    qualification.bold = True
    qualification.font.size = Pt(10.5)
    qualification.font.color.rgb = accent
    if item.institution:
        paragraph.add_run(f"  |  {item.institution}").bold = True
    meta = " · ".join(value for value in (item.dates, item.location) if value)
    if meta:
        meta_paragraph = document.add_paragraph(meta)
        meta_paragraph.paragraph_format.space_after = Pt(3)
        meta_paragraph.runs[0].italic = True
        meta_paragraph.runs[0].font.size = Pt(8.5)
        meta_paragraph.runs[0].font.color.rgb = RGBColor(95, 99, 110)
    if item.details:
        document.add_paragraph(item.details)


def _add_text_section(document: Document, title: str, value: str, accent: RGBColor, template: str) -> None:
    if not value:
        return
    _add_section_heading(document, title, accent, template)
    paragraph = document.add_paragraph(value)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08


def build_cv_document(profile: CVProfile, template: str) -> BytesIO:
    if template not in TEMPLATES:
        raise CVBuildError("Template necunoscut.")

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.65 if template != "europass" else 1.45)
    section.right_margin = Cm(1.65 if template != "europass" else 1.45)

    colors = {
        "europass": RGBColor(32, 88, 135),
        "ats": RGBColor(35, 38, 43),
        "modern": RGBColor(111, 63, 214),
    }
    accent = colors[template]
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(3)

    if template == "europass":
        banner = document.add_table(rows=1, cols=2)
        banner.autofit = False
        banner.columns[0].width = Cm(4.2)
        banner.columns[1].width = Cm(13.2)
        left, right = banner.rows[0].cells
        _set_cell_shading(left, "205887")
        _set_cell_shading(right, "EEF4F8")
        _set_cell_margins(left, 180, 180, 180, 180)
        _set_cell_margins(right, 180, 220, 180, 220)
        label = left.paragraphs[0].add_run("CURRICULUM\nVITAE")
        label.bold = True
        label.font.color.rgb = RGBColor(255, 255, 255)
        label.font.size = Pt(13)
        name = right.paragraphs[0].add_run(profile.full_name)
        name.bold = True
        name.font.size = Pt(24)
        name.font.color.rgb = accent
        if profile.headline:
            right.add_paragraph(profile.headline)
        note = document.add_paragraph("EUROPASS-INSPIRED · independent local template")
        note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        note.runs[0].font.size = Pt(7)
        note.runs[0].font.color.rgb = RGBColor(125, 132, 142)
    else:
        if template == "modern":
            accent_line = document.add_paragraph()
            _add_rule(accent_line, "6F3FD6", 22)
        name = document.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name.paragraph_format.space_after = Pt(2)
        run = name.add_run(profile.full_name)
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(26 if template == "modern" else 23)
        run.font.color.rgb = accent
        if profile.headline:
            headline = document.add_paragraph(profile.headline)
            headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
            headline.paragraph_format.space_after = Pt(4)
            headline.runs[0].font.size = Pt(10.5)

    _add_contact_line(document, profile, accent)
    _add_text_section(document, "Profil profesional", profile.summary, accent, template)

    if profile.experiences:
        _add_section_heading(document, "Experiență profesională", accent, template)
        for item in profile.experiences:
            _add_experience(document, item, accent)

    if profile.education:
        _add_section_heading(document, "Educație și formare", accent, template)
        for item in profile.education:
            _add_education(document, item, accent)

    _add_text_section(document, "Competențe", profile.skills, accent, template)
    _add_text_section(document, "Competențe digitale", profile.digital_skills, accent, template)
    _add_text_section(document, "Competențe lingvistice", profile.languages, accent, template)
    _add_text_section(document, "Proiecte", profile.projects, accent, template)
    _add_text_section(document, "Certificări", profile.certifications, accent, template)
    _add_text_section(document, "Informații suplimentare", profile.additional, accent, template)

    properties = document.core_properties
    properties.title = f"Curriculum Vitae — {profile.full_name}"
    properties.subject = "Curriculum Vitae generated locally with MORPH WRLD CV Studio"
    properties.author = profile.full_name
    properties.comments = "Generated locally. Europass-inspired is not an official Europass document."

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
