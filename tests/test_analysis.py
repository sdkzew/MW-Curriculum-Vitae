import tempfile
import unittest
from pathlib import Path

import docx

from analyzer.ats_scorer import score_cv
from analyzer.extractor import DocumentExtractionError, extract_text
from analyzer.matcher import match_job
from analyzer.suggestions import generate_suggestions


class ATSScorerTests(unittest.TestCase):
    def test_complete_cv_can_reach_one_hundred(self):
        headings = (
            "Contact test@example.com +40 712 345 678 LinkedIn\n"
            "Experience\nEducation\nSkills\nSummary\nProjects\nCertifications\n"
        )
        action_words = (
            "Developed implemented managed led created designed improved increased. "
        )
        achievements = "10 users 20 projects 30 percent 40 clients 50 releases. "
        filler = "professional engineering collaboration delivery " * 75

        result = score_cv(headings + action_words + achievements + filler)

        self.assertEqual(result.score, 100)
        self.assertEqual(result.breakdown["Sections"], (30, 30))

    def test_keywords_do_not_match_inside_other_words(self):
        result = score_cv("Knowledge is carefully profiled.\n")

        self.assertEqual(result.breakdown["Action Verbs"][0], 0)
        self.assertNotIn("summary", result.found_sections)

    def test_feedback_can_be_generated_in_english(self):
        result = score_cv("Short CV\n", "en")

        self.assertIn("Add your email", result.feedback[0])
        self.assertTrue(all("CV-ul" not in message for message in result.feedback))


class MatcherTests(unittest.TestCase):
    def test_preserves_common_technology_names(self):
        result = match_job(
            "Built services using C++, C#, .NET and Node.js.",
            "C++ C# .NET Node.js",
        )

        self.assertEqual(result.score, 100)
        self.assertTrue({"c++", "c#", ".net", "node.js"}.issubset(result.matched_keywords))

    def test_match_feedback_can_be_generated_in_english(self):
        result = match_job("Python", "Python Django Kubernetes", "en")

        self.assertEqual(result.feedback, "Average match. Your CV is missing several keywords from the job description.")


class SuggestionsTests(unittest.TestCase):
    def test_suggestions_can_be_generated_in_english(self):
        result = generate_suggestions("Short CV", 20, 0, ["python"], "en")

        self.assertEqual(result[0]["category"], "ATS compatibility")
        self.assertTrue(all("Adaugă" not in item["text"] for item in result))


class ExtractorTests(unittest.TestCase):
    def test_docx_extracts_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cv.docx"
            document = docx.Document()
            document.add_paragraph("Professional summary")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Python"
            table.cell(0, 1).text = "Five years"
            document.save(path)

            text = extract_text(str(path))

        self.assertIn("Professional summary", text)
        self.assertIn("Python", text)
        self.assertIn("Five years", text)

    def test_invalid_docx_has_stable_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.docx"
            path.write_bytes(b"not a docx")

            with self.assertRaises(DocumentExtractionError):
                extract_text(str(path))


if __name__ == "__main__":
    unittest.main()
